"""Bounded CSV/XLS/XLSX intake and owner-scoped metadata access."""

import asyncio
import csv
import hashlib
import io
import json
import os
import re
import shutil
import tempfile
import threading
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import UUID, uuid4

import pyarrow as pa
import pyarrow.parquet as pq
from fastapi import APIRouter, Depends, HTTPException, Request
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from python_multipart import MultipartParser
from python_multipart.multipart import parse_options_header
from sqlalchemy import select
from sqlalchemy.orm import Session as DbSession
from xlrd import XLRDError, open_workbook

from app.auth import current_user
from app.database import get_session
from app.models import Analysis, Conversation, Dataset, User

router = APIRouter(prefix="/api/v1/datasets")
FILE_LIMIT = 64 * 1024 * 1024
XLSX_EXPANDED_LIMIT = 64 * 1024 * 1024
FILE_COUNT = 10
ROW_LIMIT = 1100000
COLUMN_LIMIT = 50
PREVIEW_LIMIT = 50
_duplicate_upload_lock = threading.Lock()


class IntakeError(ValueError):
    pass


def inspect_rows(headers, rows):
    if len(headers) > COLUMN_LIMIT:
        raise IntakeError("Spreadsheet exceeds 50 columns.")
    if any(not h.strip() or any(ord(c) < 32 or ord(c) == 127 for c in h) for h in headers):
        raise IntakeError("Headers must be nonblank and contain no control characters.")
    if len(headers) != len(set(headers)):
        raise IntakeError("Headers must be unique.")
    if not rows:
        raise IntakeError("File needs at least one data row.")
    duplicates = Counter(tuple(row) for row in rows)
    columns = []
    arrays = {}
    warnings = []
    for index, header in enumerate(headers):
        values = [row[index] for row in rows]
        present = [value for value in values if value != ""]
        kind = "text"
        if present and all(value.lower() in ("true", "false") for value in present):
            kind = "boolean"
            converted = [None if v == "" else v.lower() == "true" for v in values]
            arrays[f"c{index + 1}"] = pa.array(converted, type=pa.bool_())
        elif (
            present
            and all(v.isascii() and v.lstrip("-").isdigit() for v in present)
            and not any(len(v.lstrip("-")) > 1 and v.lstrip("-").startswith("0") for v in present)
        ):
            try:
                converted = [None if v == "" else int(v) for v in values]
                arrays[f"c{index + 1}"] = pa.array(converted, type=pa.int64())
                kind = "integer"
            except (ValueError, OverflowError):
                pass
        elif present and all(re.fullmatch(r"-?(?:0|[1-9][0-9]*)\.[0-9]+", v) for v in present):
            try:
                scale = max(len(v.split(".")[1]) for v in present)
                if scale <= 6 and all(len(v.lstrip("-").replace(".", "")) <= 38 for v in present):
                    arrays[f"c{index + 1}"] = pa.array(
                        [None if v == "" else Decimal(v) for v in values],
                        type=pa.decimal128(38, scale),
                    )
                    kind = "decimal"
            except (InvalidOperation, ValueError):
                pass
        if f"c{index + 1}" not in arrays:
            arrays[f"c{index + 1}"] = pa.array([None if v == "" else v for v in values], type=pa.string())
            if any(len(v) > 1 and v.startswith("0") and v.isdigit() for v in present):
                warnings.append(f"{header}: leading-zero values kept as text.")
            if any("/" in v or "-" in v for v in present):
                warnings.append(f"{header}: date-like values kept as text until confirmed.")
        missing = len(values) - len(present)
        columns.append(
            {
                "id": f"c{index + 1}",
                "name": header,
                "type": kind,
                "missing_count": missing,
                "missing_rate": missing / len(rows),
            }
        )
    profile = {
        "row_count": len(rows),
        "column_count": len(headers),
        "duplicate_row_count": sum(n - 1 for n in duplicates.values()),
        "columns": columns,
        "warnings": warnings,
        "null_rule": "Only empty CSV or spreadsheet cells are null; literal text is preserved.",
    }
    return pa.table(arrays), profile


def inspect_csv(raw: bytes):
    if not raw:
        raise IntakeError("File is empty.")
    try:
        decoded = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise IntakeError("Use UTF-8 encoded CSV.") from None
    try:
        reader = csv.reader(io.StringIO(decoded, newline=""), strict=True)
        headers = next(reader, None)
        if not headers:
            raise IntakeError("CSV needs a header and at least one data row.")
        rows = []
        for row in reader:
            if len(row) != len(headers):
                raise IntakeError(f"Record at line {reader.line_num} has {len(row)} fields; expected {len(headers)}.")
            rows.append(row)
            if len(rows) > ROW_LIMIT:
                raise IntakeError("CSV exceeds 1,100,000 data rows.")
    except csv.Error:
        raise IntakeError("Malformed CSV record or quoting.") from None
    return inspect_rows(headers, rows)


def inspect_xlsx(raw: bytes):
    if not raw:
        raise IntakeError("File is empty.")
    try:
        # XLSX is compressed; bound parser input before shared strings/XML are loaded.
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > XLSX_EXPANDED_LIMIT:
                raise IntakeError("XLSX expanded contents exceed 64 MiB.")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except IntakeError:
        raise
    except (OSError, ValueError, KeyError, TypeError, InvalidFileException, zipfile.BadZipFile):
        raise IntakeError("Malformed or unsupported XLSX workbook.") from None
    try:
        worksheet = next((sheet for sheet in workbook.worksheets if sheet.max_row and sheet.max_column), None)
        if worksheet is None:
            raise IntakeError("XLSX needs a worksheet with a header and at least one data row.")
        if worksheet.max_column > COLUMN_LIMIT:
            raise IntakeError("XLSX exceeds 50 columns.")
        headers = None
        rows = []
        for raw_row in worksheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in raw_row]
            if not any(values):
                continue
            if headers is None:
                headers = values
                continue
            if len(rows) >= ROW_LIMIT:
                raise IntakeError("XLSX exceeds 1,100,000 data rows.")
            rows.append(values + [""] * (len(headers) - len(values)))
        if headers is None:
            raise IntakeError("XLSX needs a header and at least one data row.")
        if any(len(row) > len(headers) and any(row[len(headers) :]) for row in rows):
            raise IntakeError("XLSX contains data beyond its header columns.")
        return inspect_rows(headers, [row[: len(headers)] for row in rows])
    finally:
        workbook.close()


def inspect_xls(raw: bytes):
    if not raw:
        raise IntakeError("File is empty.")
    try:
        workbook = open_workbook(file_contents=raw, on_demand=True)
    except (OSError, XLRDError, ValueError):
        raise IntakeError("Malformed or unsupported XLS workbook.") from None
    try:
        worksheet = next((sheet for sheet in workbook.sheets() if sheet.nrows and sheet.ncols), None)
        if worksheet is None:
            raise IntakeError("XLS needs a worksheet with a header and at least one data row.")
        if worksheet.ncols > COLUMN_LIMIT:
            raise IntakeError("XLS exceeds 50 columns.")
        rows = []
        headers = None
        for row_index in range(worksheet.nrows):
            values = ["" if value is None else str(value) for value in worksheet.row_values(row_index)]
            if not any(values):
                continue
            if headers is None:
                headers = values
                continue
            if len(rows) >= ROW_LIMIT:
                raise IntakeError("XLS exceeds 1,100,000 data rows.")
            rows.append(values)
        if headers is None:
            raise IntakeError("XLS needs a header and at least one data row.")
        return inspect_rows(headers, [row[: len(headers)] for row in rows])
    finally:
        workbook.release_resources()


def public_dataset(dataset: Dataset):
    profile = json.loads(dataset.profile_json)
    return {
        "id": str(dataset.id),
        "display_name": dataset.display_name,
        "version": dataset.version,
        "status": dataset.status,
        "content_hash": dataset.content_hash,
        "created_at": dataset.created_at.isoformat(),
        "profile": profile,
    }


def owned(db: DbSession, user: User, dataset_id: UUID) -> Dataset:
    dataset = db.scalar(select(Dataset).where(Dataset.id == dataset_id, Dataset.owner_id == user.id))
    if not dataset:
        raise HTTPException(404)
    return dataset


@router.get("")
def list_datasets(request: Request, user: User = Depends(current_user), db: DbSession = Depends(get_session)):
    items = db.scalars(select(Dataset).where(Dataset.owner_id == user.id).order_by(Dataset.created_at.desc())).all()
    return {
        "request_id": request.state.request_id,
        "data": [public_dataset(item) for item in items],
    }


@router.get("/{dataset_id}")
def detail(
    dataset_id: UUID,
    request: Request,
    user: User = Depends(current_user),
    db: DbSession = Depends(get_session),
):
    return {
        "request_id": request.state.request_id,
        "data": public_dataset(owned(db, user, dataset_id)),
    }


@router.get("/{dataset_id}/preview")
def preview(
    dataset_id: UUID,
    request: Request,
    user: User = Depends(current_user),
    db: DbSession = Depends(get_session),
):
    dataset = owned(db, user, dataset_id)
    parquet = pq.ParquetFile(request.app.state.settings.data_dir / dataset.parquet_path)
    columns = json.loads(dataset.schema_json)
    try:
        rows = [
            {key: str(value) if isinstance(value, Decimal) else value for key, value in row.items()}
            for row in next(parquet.iter_batches(batch_size=PREVIEW_LIMIT)).to_pylist()
        ]
    finally:
        parquet.close()
    return {
        "request_id": request.state.request_id,
        "data": {
            "columns": columns,
            "rows": rows,
            "returned_rows": len(rows),
            "total_rows": dataset.row_count,
            "truncated": dataset.row_count > PREVIEW_LIMIT,
        },
    }


async def receive_files(request: Request, root: Path):
    content_type, params = parse_options_header(request.headers.get("content-type", ""))
    if content_type != b"multipart/form-data" or b"boundary" not in params:
        raise HTTPException(422)
    files = []
    finished = False
    part = {"field": b"", "value": b"", "headers": {}, "name": "", "file": None, "size": 0, "error": None}

    def begin():
        part.update(field=b"", value=b"", headers={}, name="", file=None, size=0, error=None)

    def header_field(data, start, end):
        part["field"] += data[start:end]

    def header_value(data, start, end):
        part["value"] += data[start:end]

    def header_end():
        part["headers"][part["field"].lower()] = part["value"]
        part["field"] = part["value"] = b""

    def headers_finished():
        _, options = parse_options_header(part["headers"].get(b"content-disposition", b""))
        if options.get(b"name") != b"files" or b"filename" not in options:
            raise IntakeError("Expected file fields named files.")
        if len(files) >= FILE_COUNT:
            raise IntakeError("Choose at most 10 files.")
        part["name"] = options[b"filename"].decode("utf-8", errors="replace")
        part["file"] = tempfile.NamedTemporaryFile(dir=root, delete=False)

    def part_data(data, start, end):
        part["size"] += end - start
        if part["size"] > FILE_LIMIT:
            part["error"] = "File exceeds 64 MiB."
        elif not part["error"]:
            part["file"].write(data[start:end])

    def part_end():
        part["file"].close()
        files.append((part["name"], Path(part["file"].name), part["error"]))

    def end():
        nonlocal finished
        finished = True

    parser = MultipartParser(
        params[b"boundary"],
        {
            "on_part_begin": begin,
            "on_header_field": header_field,
            "on_header_value": header_value,
            "on_header_end": header_end,
            "on_headers_finished": headers_finished,
            "on_part_data": part_data,
            "on_part_end": part_end,
            "on_end": end,
        },
        max_size=256 * 1024 * 1024,
    )
    try:
        async for chunk in request.stream():
            parser.write(chunk)
        parser.finalize()
        if not finished or not files:
            raise IntakeError("Upload is incomplete or contains no files.")
        return files
    except BaseException:
        if part.get("file") and not part["file"].closed:
            part["file"].close()
            Path(part["file"].name).unlink(missing_ok=True)
        for _, path, _ in files:
            path.unlink(missing_ok=True)
        raise


def process_file(filename, incoming, incoming_error, owner_id, root, factory):
    dataset_id = uuid4()
    work = Path(tempfile.mkdtemp(prefix="intake-", dir=root))
    final = root / str(dataset_id)
    try:
        if incoming_error:
            raise IntakeError(incoming_error)
        raw = incoming.read_bytes()
        content_hash = hashlib.sha256(raw).hexdigest()
        suffix = Path(filename or "").suffix.lower()
        if suffix == ".csv":
            table, profile = inspect_csv(raw)
        elif suffix == ".xlsx":
            table, profile = inspect_xlsx(raw)
        elif suffix == ".xls":
            table, profile = inspect_xls(raw)
        else:
            raise IntakeError("Only .csv, .xls, and .xlsx files are supported.")
        original_name = f"original{suffix}"
        with _duplicate_upload_lock:
            with factory() as db:
                duplicate = db.scalar(
                    select(Dataset.id)
                    .where(Dataset.owner_id == owner_id, Dataset.content_hash == content_hash)
                    .limit(1)
                )
            if duplicate:
                raise IntakeError("This dataset has already been uploaded.")
            (work / original_name).write_bytes(raw)
            pq.write_table(table, work / "snapshot.parquet")
            os.replace(work, final)
            dataset = Dataset(
                id=dataset_id,
                owner_id=owner_id,
                display_name=Path(filename or "upload.csv").name[:255],
                content_hash=content_hash,
                version=1,
                status="ready",
                original_path=f"datasets/{dataset_id}/{original_name}",
                parquet_path=f"datasets/{dataset_id}/snapshot.parquet",
                schema_json=json.dumps(profile["columns"]),
                profile_json=json.dumps(profile),
                row_count=profile["row_count"],
            )
            with factory() as db:
                db.add(dataset)
                db.commit()
        return {"filename": filename, "status": "accepted", "dataset": public_dataset(dataset)}
    except IntakeError as exc:
        return {"filename": filename, "status": "rejected", "error": str(exc)}
    except Exception:
        shutil.rmtree(final, ignore_errors=True)
        return {"filename": filename, "status": "rejected", "error": "Storage failed. Try again."}
    finally:
        shutil.rmtree(work, ignore_errors=True)


@router.post("")
async def upload(request: Request, user: User = Depends(current_user)):
    from time import perf_counter

    from app.observability import emit

    operation_started = perf_counter()
    root = request.app.state.settings.data_dir / "datasets"
    root.mkdir(parents=True, exist_ok=True)
    try:
        files = await receive_files(request, root)
    except IntakeError as exc:
        raise HTTPException(413, str(exc)) from None
    results = []
    try:
        for filename, incoming, incoming_error in files:
            task = asyncio.create_task(
                asyncio.to_thread(
                    process_file, filename, incoming, incoming_error, user.id, root, request.app.state.sessions
                )
            )
            try:
                result = await asyncio.shield(task)
            except asyncio.CancelledError:
                await task
                raise
            results.append(result)
    finally:
        for _, incoming, _ in files:
            incoming.unlink(missing_ok=True)
        emit("ingestion", duration_ms=round((perf_counter() - operation_started) * 1000))
    return {"request_id": request.state.request_id, "data": results}


@router.delete("/{dataset_id}")
def delete(
    dataset_id: UUID,
    request: Request,
    user: User = Depends(current_user),
    db: DbSession = Depends(get_session),
):
    dataset = owned(db, user, dataset_id)
    directory = request.app.state.settings.data_dir / "datasets" / str(dataset.id)
    conversations = db.scalars(select(Conversation).where(Conversation.owner_id == user.id)).all()
    dependent_conversations = {
        analysis.conversation_id
        for analysis in db.scalars(select(Analysis).where(Analysis.owner_id == user.id)).all()
        if any(entry["id"] == str(dataset.id) for entry in json.loads(analysis.dataset_versions_json))
    }
    for conversation in conversations:
        versions = json.loads(conversation.selected_dataset_versions_json)
        if conversation.id in dependent_conversations or any(entry["id"] == str(dataset.id) for entry in versions):
            db.delete(conversation)
    db.delete(dataset)
    db.commit()
    request.app.state.execution_cache.invalidate(user.id, dataset.id)
    from app.observability import emit

    emit("cache", reason="source_deleted")
    shutil.rmtree(directory, ignore_errors=True)
    return {"request_id": request.state.request_id, "data": {"deleted": True}}
